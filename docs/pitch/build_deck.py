from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
import sys

MESH = sys.argv[1]; SHOT = sys.argv[2]; OUT = sys.argv[3]
BG = RGBColor(0x0E, 0x14, 0x26); CARD = RGBColor(0x1A, 0x22, 0x3C); ACC = RGBColor(0xFF, 0x5A, 0x4F)
WHITE = RGBColor(255, 255, 255); MUTED = RGBColor(0xA9, 0xB3, 0xCC); GREEN = RGBColor(0x4C, 0xD9, 0x8B); AMBER = RGBColor(0xFF, 0xB8, 0x4D)

p = Presentation(); p.slide_width = Inches(13.333); p.slide_height = Inches(7.5)
blank = p.slide_layouts[6]


def text(s, x, y, w, h, t, size=16, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = t; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = "Calibri"
    tf.paragraphs[0].alignment = align
    return tb


def slide(n, kicker, title):
    s = p.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    text(s, 0.6, 0.35, 8, 0.4, kicker.upper(), 13, ACC, True)
    text(s, 0.6, 0.7, 12.1, 1.0, title, 34, WHITE, True)
    text(s, 12.2, 7.0, 0.8, 0.3, f"{n}/6", 11, MUTED, align=PP_ALIGN.RIGHT)
    return s


def card(s, x, y, w, h, head, body, accent=ACC, hs=17, bs=13):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.adjustments[0] = 0.08; b.fill.solid(); b.fill.fore_color.rgb = CARD; b.line.fill.background()
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.18), Inches(0.07), Inches(0.4)); bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    text(s, x + 0.2, y + 0.12, w - 0.35, 0.5, head, hs, WHITE, True)
    text(s, x + 0.2, y + 0.62, w - 0.35, h - 0.7, body, bs, MUTED)


def bullets(s, x, y, w, items, size=16, gap=0.62):
    for i, t in enumerate(items):
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y + i * gap + 0.11), Inches(0.13), Inches(0.13)); dot.fill.solid(); dot.fill.fore_color.rgb = ACC; dot.line.fill.background()
        text(s, x + 0.25, y + i * gap, w - 0.25, gap, t, size, WHITE)


# 1 story
s = slide(1, "KHOJ", "I got lost in a mela. My phone was useless.")
text(s, 0.6, 1.75, 6.3, 0.5, "How this idea started", 20, ACC, True)
bullets(s, 0.6, 2.3, 6.3, ["I was in a huge crowd and got separated from my family.",
                           "No signal. Calls and WhatsApp did not go through.",
                           "The only hope was a loudspeaker announcement.",
                           "But every phone around me still had Bluetooth. Why can't phones just pass the message along?"], 17, 0.95)
card(s, 7.4, 1.85, 5.3, 2.1, "The same story, again and again", "Maha Kumbh 2025: tens of thousands of people were lost. Concerts, stadiums and railway platforms are the same: a crowd gets big, the network collapses, and a loudspeaker is the only tool left.", ACC, 18, 14)
card(s, 7.4, 4.2, 5.3, 2.1, "So we built KHOJ", "Find your family in a crowd with no internet and no SIM. Every phone running KHOJ becomes a relay, so the crowd itself becomes the network.", GREEN, 18, 14)

# 2 features
s = slide(2, "What KHOJ does", "Main features")
F = [("Live radar", "Each family member is a bubble at their real direction and distance."),
     ("Find", "An arrow that says how it knows. Shows nothing rather than a guess it cannot support."),
     ("Family pass", "Join a group with one QR code. Set how far someone may wander before everyone is alerted."),
     ("Wristband QR", "Any phone camera reads it with no app and no internet. Shows a Call button for the child's guardian."),
     ("Light up", "Make a family member's torch, screen and sound go off so you can spot them in a crowd (they must allow it)."),
     ("Voice assistant", "Speak in English, Hindi or Tamil. It runs on the phone and handles the whole app."),
     ("SOS and fall alert", "Hold or shake to alert the group. Fall detection sends an alert automatically."),
     ("Simple mode", "A big-button screen for elders and first-time users.")]
for i, (h, b) in enumerate(F):
    card(s, 0.6 + (i % 4) * 3.1, 1.8 + (i // 4) * 2.6, 2.95, 2.4, h, b, ACC if i % 2 == 0 else AMBER, 17, 13)
text(s, 0.6, 7.0, 10, 0.3, "Android today. Calling a wristband number needs ordinary phone signal; nothing else needs internet.", 11, MUTED)

# 3 technical
s = slide(3, "Technical", "Offline ways phones talk to each other")
T = [("Bluetooth LE mesh", "BLE 5 extended advertising. Managed flooding with relay delay and duplicate suppression, store-carry-forward, ACK retries.", ACC),
     ("Google Nearby Connections", "Cluster of phones. Starts on Bluetooth, upgrades to Wi-Fi LAN. Verified between our phones.", AMBER),
     ("Wi-Fi Direct and Wi-Fi Aware", "Higher-speed links when the phones support them. Wi-Fi LAN was the verified upgrade; Direct is not yet verified.", AMBER),
     ("GATT link", "Direct Bluetooth connection as a fallback bearer.", ACC),
     ("Optional internet relay", "If a phone happens to have internet it can bridge groups. Never required.", MUTED)]
for i, (h, b, c) in enumerate(T):
    card(s, 0.6 + (i % 3) * 4.15, 1.8 + (i // 3) * 2.45, 3.95, 2.25, h, b, c, 17, 13)
card(s, 8.9, 4.25, 3.95, 2.25, "Private by design", "Packets are sealed with AES-256-GCM using a key only the family holds. Strangers can relay but not read. Identifiers rotate every 15 minutes.", GREEN, 17, 13)
text(s, 0.6, 6.8, 11, 0.35, "Tested on 3 real phones. Range and crowd-scale behaviour still need more field testing.", 12, MUTED)

# 4 innovation
s = slide(4, "Our innovation", "The mesh heals itself: packet regeneration")
def arrow(x, y):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(0.5), Inches(0.5)); a.fill.solid(); a.fill.fore_color.rgb = ACC; a.line.fill.background()
card(s, 0.6, 1.8, 3.9, 2.9, "1. What the phone already knows", "Last known location. Timestamp and how old it is. Speed and motion (walking, still). Signal-strength history. Copies of the packet relayed by neighbours.", AMBER, 17, 13)
arrow(4.6, 2.9)
card(s, 5.2, 1.8, 3.9, 2.9, "2. On-device AI on the NPU", "A small model trained on a laptop, exported to LiteRT and run on the phone's Qualcomm NPU through the QNN delegate. Fast, private, no internet.", ACC, 17, 13)
arrow(9.2, 2.9)
card(s, 9.8, 1.8, 2.95, 2.9, "3. Regenerated packet", "The missing update is rebuilt and tagged regenerated, with a confidence. A real packet replaces it the moment one arrives.", GREEN, 17, 13)
card(s, 0.6, 4.95, 4.0, 1.85, "Why it matters", "In a crowd, packets drop in bursts. Instead of a family member vanishing from the radar, their last good data fills the gap.", GREEN, 16, 13)
card(s, 4.75, 4.95, 4.0, 1.85, "Safety rule", "A regenerated position is never shown as live. If the gap passes 20 seconds, KHOJ shows nothing rather than guess.", ACC, 16, 13)
card(s, 8.9, 4.95, 3.85, 1.85, "Status, honestly", "Design and prototype. Reed-Solomon recovery in direct sharing: 92.8% vs 54.6% (simulation). No measured regeneration results yet. It does not extend range.", AMBER, 16, 12)

# 5 mesh
s = slide(5, "The mesh", "The crowd becomes the network")
s.shapes.add_picture(MESH, Inches(0.5), Inches(1.7), width=Inches(6.9))
text(s, 0.5, 6.45, 7.6, 0.5, "Illustration: a message hops phone to phone from sender A to receiver O. No tower, no internet.", 12, MUTED)
bullets(s, 8.4, 1.8, 4.5, ["More phones means more routes, so the mesh gets stronger in a crowd.",
                           "SOS and Find messages get priority over everything else.",
                           "Phones that are out of range carry the message and pass it on later.",
                           "Find is honest about age: a live position is used as it is, a 10 to 20 second old one is predicted, and nothing is shown after 20 seconds."], 14, 1.15)

# 6 conclusion
s = slide(6, "Conclusion", "Losing signal should not mean losing your people")
card(s, 0.6, 1.8, 4.0, 2.4, "The problem", "In big crowds the network fails exactly when families need it most, and a loudspeaker is all that is left.", ACC, 18, 14)
card(s, 4.75, 1.8, 4.0, 2.4, "Our answer", "KHOJ: phone-to-phone mesh, a radar and arrow to find family, a wristband any camera can read, and voice in Hindi and Tamil.", GREEN, 18, 14)
card(s, 8.9, 1.8, 3.85, 2.4, "Where we are", "Working app on 3 real phones. Next: crowd-scale field tests and in-mesh packet regeneration on the NPU.", AMBER, 18, 14)
text(s, 0.6, 4.6, 8, 0.5, "Join the beta", 24, ACC, True)
text(s, 0.6, 5.2, 7.5, 1.2, "Early software. Android only for now. Not a replacement for calling 112.", 15, MUTED)
s.shapes.add_picture(SHOT, Inches(9.2), Inches(4.45), height=Inches(2.6))
p.save(OUT); print("saved", OUT)
